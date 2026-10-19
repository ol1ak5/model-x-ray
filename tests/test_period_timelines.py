from datetime import datetime
from pathlib import Path

import openpyxl

from model_x_ray.agents.locator import check_target, prior_cell
from model_x_ray.rules import is_period_cell, scan_model
from model_x_ray.spreadsheet.xlsx import load_xlsx


def write_timelines(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    valuation = workbook.active
    valuation.title = "Valuation output"
    valuation.append(["", "Base year", 1, 2, 3, 4, 5, "Terminal year"])
    valuation.append([])
    valuation.append(["Revenues", 100, "=B3*1.1", "=C3*1.1", 133.1, "=E3*1.1", "=F3*1.1", "=G3*1.02"])
    valuation.append(["EBIT", "=B3*0.2", "=C3*0.2", "=D3*0.2", "=E3*0.2", "=F3*0.2", "=G3*0.2", "=H3*0.2"])
    monthly = workbook.create_sheet("Acc")
    monthly["B2"], monthly["C2"] = datetime(2021, 12, 31), "=EOMONTH(B2,12)"
    monthly["D2"], monthly["E2"] = "=EOMONTH(C2,12)", "=EOMONTH(D2,12)"
    for column in "BCDE":
        monthly[f"{column}2"].number_format = "dd/mm/yyyy"
    monthly["A4"], monthly["B4"] = "Toll revenue", 50
    for previous, column in zip("BCD", "CDE"):
        monthly[f"{column}4"] = f"={previous}4*1.03"
    workbook.save(path)
    return path


def test_relative_years_and_date_headers_are_read_as_timelines(tmp_path: Path) -> None:
    book = load_xlsx(write_timelines(tmp_path / "timelines.xlsx"))

    assert [book.column_header(f"Valuation output!{c}3") for c in "BCDEFGH"] == ["Base year", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5", "Terminal year"]
    assert is_period_cell(book, "Valuation output!B3")
    assert prior_cell(book, "Valuation output!C3") == "Valuation output!B3"
    assert [o["cell"] for o in scan_model(book)] == ["Valuation output!E3"]
    target, question = check_target(book)
    assert target in {"Valuation output!F3", "Valuation output!E4"}
    assert "Year" in question and "Valuation output!E3" in question

    assert [book.column_header(f"Acc!{c}4") for c in "BCDE"] == ["2021-12-31", "2022-12-31", "2023-12-31", "2024-12-31"]
    assert prior_cell(book, "Acc!D4") == "Acc!C4"


def test_a_numbered_list_is_not_mistaken_for_a_timeline(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Deals"
    sheet.append(["Deal", "Volume", "Price"])
    for number in range(1, 6):
        sheet.append([number, 100 * number, "=B{0}*2".format(number + 1)])
    workbook.save(tmp_path / "deals.xlsx")
    book = load_xlsx(tmp_path / "deals.xlsx")

    assert book.column_header("Deals!B3") == "Volume"
