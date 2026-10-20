from pathlib import Path
from uuid import uuid4

import openpyxl

from model_x_ray.benchmark.runner import Setup, build_setup_graph, drive_graph
from model_x_ray.rules import isolated_break, rank_outliers, scan_model
from model_x_ray.spreadsheet.sandbox import XlsxSandbox
from model_x_ray.spreadsheet.source import XlsxSource
from model_x_ray.spreadsheet.xlsx import load_xlsx

YEARS = "DEFGHIJK"


def write_transition_model(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    model = workbook.active
    model.title = "Valuation"
    inputs = workbook.create_sheet("Inputs")
    inputs["A1"], inputs["B1"] = "High growth", 0.12
    inputs["A2"], inputs["B2"] = "Stable growth", 0.03
    inputs["A3"], inputs["B3"] = "Opex share", 0.4
    model["A3"] = "EUR m"
    for index, column in enumerate(YEARS):
        model[f"{column}3"] = 2026 + index
    model["A5"], model["A6"], model["A7"], model["A8"] = "Growth", "Revenue", "Opex", "EBITDA"
    model["C6"] = 500
    for index, column in enumerate(YEARS):
        previous = chr(ord(column) - 1)
        step = index - 3
        if index < 4:
            model[f"{column}5"] = "=Inputs!$B$1"
        elif index < 7:
            model[f"{column}5"] = f"=Inputs!$B$1-(Inputs!$B$1-Inputs!$B$2)*{step}/4"
        else:
            model[f"{column}5"] = "=Inputs!$B$2"
        model[f"{column}6"] = f"={previous}6*(1+{column}5)"
        model[f"{column}7"] = f"={column}6*Inputs!$B$3"
        model[f"{column}8"] = f"={column}6-{column}7"
    model["G7"] = "=G5*Inputs!$B$3"
    workbook.save(path)
    return path


def test_an_isolated_break_is_ranked_before_a_deliberate_block_of_transition_years(tmp_path: Path) -> None:
    book = load_xlsx(write_transition_model(tmp_path / "valuation.xlsx"))

    flagged = [o["cell"] for o in scan_model(book)]
    ranked = [o["cell"] for o in rank_outliers(book, scan_model(book))]

    assert {"Valuation!H5", "Valuation!I5", "Valuation!G7"} <= set(flagged)
    assert isolated_break(book, "Valuation!G7")
    assert not isolated_break(book, "Valuation!H5")
    assert len(book.all_dependents("Valuation!H5")) > len(book.all_dependents("Valuation!G7"))
    assert ranked[0] == "Valuation!G7"


def test_the_isolated_break_is_the_first_candidate_and_gets_fixed(tmp_path: Path) -> None:
    from conftest import ScriptedVerdict
    from test_investigator import ScriptedDiagnosis

    path = write_transition_model(tmp_path / "valuation.xlsx")
    investigator = ScriptedDiagnosis(turns=[([], None)])
    setup = Setup(name="scripted", kind="investigator", model="scripted-investigator", verifier="scripted-verifier")
    source = XlsxSource({"valuation": path})
    graph = build_setup_graph(setup, source, XlsxSandbox(source, tmp_path / "corrected"), "valuation", lambda model_id: ScriptedVerdict() if model_id == "scripted-verifier" else investigator)
    final = drive_graph(graph, {"spreadsheet_id": "valuation", "question": ""}, {"configurable": {"thread_id": uuid4().hex}})

    assert "  1. Valuation!G7:" in investigator.briefs[0]
    assert final["patch"]["cell"] == "Valuation!G7"
    assert final["patch"]["new_formula"] == "=G6*Inputs!$B$3"
    assert final["verifications"][-1]["verdict"] == "confirmed"


def test_a_clean_model_with_cells_the_engine_cannot_calculate_is_still_reviewed(tmp_path: Path) -> None:
    from conftest import ScriptedVerdict
    from test_investigator import ScriptedDiagnosis

    path = write_transition_model(tmp_path / "valuation.xlsx")
    workbook = openpyxl.load_workbook(path)
    model = workbook["Valuation"]
    model["G7"] = "=G6*Inputs!$B$3"
    for index, column in enumerate(YEARS):
        model[f"{column}5"] = "=Inputs!$B$1"
        model[f"{column}10"] = f'=IF({column}3>2030," ",{column}8/1.1)'
        model[f"{column}11"] = f"={column}10*2+{column}6+{column}7+{column}8"
    model["A10"], model["A11"] = "Discounted EBITDA", "Value"
    workbook.save(path)
    investigator = ScriptedDiagnosis(turns=[([], None)])
    setup = Setup(name="scripted", kind="investigator", model="scripted-investigator", verifier="scripted-verifier")
    source = XlsxSource({"valuation": path})
    graph = build_setup_graph(setup, source, XlsxSandbox(source, tmp_path / "corrected"), "valuation", lambda model_id: ScriptedVerdict() if model_id == "scripted-verifier" else investigator)
    final = drive_graph(graph, {"spreadsheet_id": "valuation", "question": ""}, {"configurable": {"thread_id": uuid4().hex}})

    target = f"Valuation!{final['target']['cell']}"
    assert isinstance(load_xlsx(path).value(target), float)
    assert final.get("patch") is None
