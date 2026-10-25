import base64
import io
import json
from pathlib import Path

import openpyxl
from conftest import ScriptedDiagnosis
from fastapi.testclient import TestClient
from langgraph.checkpoint.memory import InMemorySaver

from model_x_ray.api import AgentModels, UsageCap, create_app
from model_x_ray.spreadsheet.sandbox import CHANGE_LOG_SHEET, XlsxSandbox
from model_x_ray.spreadsheet.source import XlsxSource

QUESTION = "Why did EBITDA 2027 jump on the P&L sheet?"


def instance(demo_workbook_path: Path, scratch: Path, scripted_verdict) -> TestClient:
    source = XlsxSource({"demo": demo_workbook_path})
    corrected = scratch / "corrected"
    models = AgentModels(llm=ScriptedDiagnosis(turns=[(["P&L!H7"], "P&L!H7")]), verifier=scripted_verdict)
    app = create_app(
        source, XlsxSandbox(source, corrected), models, InMemorySaver(), UsageCap(per_visitor=10, total=100),
        upload_dir=scratch / "uploads", copy_dir=corrected,
    )
    return TestClient(app)


def events(text: str) -> list[tuple[str, dict]]:
    found = []
    for block in text.strip().split("\n\n"):
        fields = dict(line.split(": ", 1) for line in block.split("\n") if ": " in line)
        found.append((fields["event"], json.loads(fields["data"])))
    return found


def final_view(text: str) -> dict:
    views = [data for event, data in events(text) if event == "view"]
    assert views, text
    return views[-1]


def with_file(request: dict, content: bytes, name: str = "My model.xlsx") -> dict:
    return {"data": {"request": json.dumps(request)}, "files": {"file": (name, content, "application/octet-stream")}}


def test_a_review_started_on_one_instance_is_approved_on_another_that_never_saw_the_upload(demo_workbook_path: Path, scripted_verdict, tmp_path: Path) -> None:
    content = demo_workbook_path.read_bytes()
    first = instance(demo_workbook_path, tmp_path / "first", scripted_verdict)
    second = instance(demo_workbook_path, tmp_path / "second", scripted_verdict)

    started = first.post("/investigations/stream", **with_file({"question": QUESTION}, content))

    assert started.status_code == 200, started.text
    paused = final_view(started.text)
    assert paused["status"] == "needs_review"
    assert paused["spreadsheet_id"].startswith("upload-")
    assert paused["sheet"]["sheet"] == "P&L", "the grid travels with the view, so no second request is needed"
    assert any(cell["ref"] == "P&L!H7" for row in paused["sheet"]["rows"] for cell in row["cells"])
    assert paused["corrected_file"] is None

    carried = json.loads(json.dumps(paused))
    approved = second.post(f"/investigations/{paused['id']}/resume/stream", **with_file({"action": "approve", "reason": "typed opex", "state": carried}, content))

    assert approved.status_code == 200, approved.text
    done = final_view(approved.text)
    assert done["status"] == "done"
    corrected = done["corrected_file"]
    assert corrected["name"].endswith(".xlsx")
    book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(corrected["base64"])))
    assert book["P&L"]["H7"].value == "=G7*(1+Assumptions!$C$7)", "the download is the visitor's workbook with the approved fix in it"
    assert CHANGE_LOG_SHEET in book.sheetnames


def test_the_same_file_gets_the_same_id_everywhere_and_another_file_cannot_resume_the_review(demo_workbook_path: Path, scripted_verdict, tmp_path: Path) -> None:
    content = demo_workbook_path.read_bytes()
    first = instance(demo_workbook_path, tmp_path / "first", scripted_verdict)
    second = instance(demo_workbook_path, tmp_path / "second", scripted_verdict)

    uploaded = first.post("/spreadsheets", params={"name": "My model.xlsx"}, content=content).json()["id"]
    assert second.post("/spreadsheets", params={"name": "Renamed.xlsx"}, content=content).json()["id"] == uploaded

    paused = final_view(first.post("/investigations/stream", **with_file({"question": QUESTION}, content)).text)
    assert paused["spreadsheet_id"] == uploaded

    other = openpyxl.load_workbook(io.BytesIO(content))
    other["Assumptions"]["C7"] = 0.5
    buffer = io.BytesIO()
    other.save(buffer)

    refused = second.post(f"/investigations/{paused['id']}/resume/stream", **with_file({"action": "approve", "reason": "typed opex", "state": paused}, buffer.getvalue()))

    assert refused.status_code == 400
    assert "not the file" in refused.json()["detail"]


def test_a_request_without_a_readable_workbook_gets_a_plain_answer(demo_workbook_path: Path, scripted_verdict, tmp_path: Path) -> None:
    client = instance(demo_workbook_path, tmp_path, scripted_verdict)

    not_excel = client.post("/investigations/stream", **with_file({"question": QUESTION}, b"hello", name="notes.txt"))
    assert not_excel.status_code == 400
    assert ".xlsx" in not_excel.json()["detail"]

    broken = client.post("/investigations/stream", **with_file({"question": QUESTION}, b"not a zip", name="model.xlsx"))
    assert broken.status_code == 400
    assert "not a readable .xlsx" in broken.json()["detail"]

    bad_request = client.post("/investigations/stream", data={"request": "{not json"}, files={"file": ("model.xlsx", demo_workbook_path.read_bytes(), "application/octet-stream")})
    assert bad_request.status_code == 400


def test_behind_the_vercel_proxy_each_visitor_gets_their_own_daily_limit(demo_workbook_path: Path, scripted_verdict, tmp_path: Path) -> None:
    source = XlsxSource({"demo": demo_workbook_path})
    models = AgentModels(llm=ScriptedDiagnosis(turns=[(["P&L!H7"], "P&L!H7")] * 3), verifier=scripted_verdict)
    client = TestClient(create_app(source, XlsxSandbox(source, tmp_path / "corrected"), models, InMemorySaver(), UsageCap(per_visitor=1, total=100)))
    body = {"spreadsheet_id": "demo", "question": QUESTION}

    first = client.post("/investigations/stream", json=body, headers={"x-forwarded-for": "203.0.113.7, 10.0.0.1"})
    assert final_view(first.text)["status"] == "needs_review"

    again = client.post("/investigations/stream", json=body, headers={"x-forwarded-for": "203.0.113.7"})
    assert again.status_code == 429

    other = client.post("/investigations/stream", json=body, headers={"x-forwarded-for": "198.51.100.4"})
    assert other.status_code == 200
    assert final_view(other.text)["status"] == "needs_review"
