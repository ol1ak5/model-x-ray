from pathlib import Path

import openpyxl

from model_x_ray.spreadsheet.xlsx import load_xlsx
from model_x_ray.structure import describe, profile


def write_mixed_model(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    inputs = workbook.active
    inputs.title = "Inputs"
    inputs["A1"] = "Assumptions"
    for row, (label, value) in enumerate([("Revenue growth", 0.08), ("COGS share", 0.55), ("Opex share", 0.2), ("Tax rate", 0.25), ("Base revenue", 1000)], start=3):
        inputs[f"A{row}"], inputs[f"B{row}"] = label, value
    pnl = workbook.create_sheet("P&L")
    pnl.append(["P&L"])
    pnl.append([])
    pnl.append(["", 2025, 2026, 2027, 2028, 2029])
    pnl.append(["Revenue", "=Inputs!B7", "=B4*(1+Inputs!$B$3)", "=C4*(1+Inputs!$B$3)", "=D4*(1+Inputs!$B$3)", "=E4*(1+Inputs!$B$3)"])
    pnl.append(["COGS"] + [f"=-{c}4*Inputs!$B$4" for c in "BCDEF"])
    pnl.append(["Gross profit"] + [f"={c}4+{c}5" for c in "BCDEF"])
    pnl.append(["Opex"] + [f"=-{c}4*Inputs!$B$5" for c in "BCDEF"])
    pnl.append(["EBITDA"] + [f"={c}6+{c}7" for c in "BCDEF"])
    deals = workbook.create_sheet("Deals")
    deals.append(["Deal", "Volume", "Price", "Value"])
    for row, (volume, price) in enumerate([(100, 2.5), (80, 2.6), (120, 2.4), (90, 2.7), (60, 2.5)], start=2):
        deals.append([f"D-{row}", volume, price, f"=B{row}*C{row}"])
    deals["A8"], deals["D8"] = "Total", "=SUM(D2:D6)"
    summary = workbook.create_sheet("Summary")
    summary["A1"], summary["B1"] = "EBITDA 2029", "='P&L'!F8"
    summary["A2"], summary["B2"] = "Deal value", "=Deals!D8"
    workbook.save(path)
    return path


def test_each_sheet_is_classified_by_what_its_formulas_do(tmp_path: Path) -> None:
    book = load_xlsx(write_mixed_model(tmp_path / "mixed.xlsx"))

    blocks = profile(book)
    kinds = {(block.sheet, block.kind) for block in blocks}

    assert kinds == {("Inputs", "inputs"), ("P&L", "temporal"), ("Deals", "table"), ("Summary", "summary")}
    temporal = next(block for block in blocks if block.kind == "temporal")
    assert (temporal.header_row, temporal.axis, temporal.top, temporal.bottom) == (3, "across", 4, 8)
    assert "2025" in temporal.evidence and "2029" in temporal.evidence
    table = next(block for block in blocks if block.kind == "table")
    assert (table.header_row, table.axis, table.top, table.bottom, table.totals) == (1, "down", 2, 6, ["Deals!D8"])
    assert "5 rows" in table.evidence

    text = describe(blocks)
    assert text.startswith("Mixed model:")
    assert "time series on P&L (2025 to 2029)" in text
    assert "table on Deals (5 rows, total in D8)" in text
    assert "assumptions on Inputs" in text


def test_a_projection_alone_is_a_time_series_model(demo_workbook_path: Path) -> None:
    blocks = profile(load_xlsx(demo_workbook_path))

    assert {block.kind for block in blocks} >= {"temporal"}
    assert describe(blocks).startswith("Time-series model:")
