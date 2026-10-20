from pathlib import Path
from uuid import uuid4

import openpyxl
import pytest

from conftest import ScriptedVerdict
from test_investigator import ScriptedDiagnosis

from model_x_ray.benchmark.runner import Setup, build_setup_graph, drive_graph
from model_x_ray.spreadsheet.engine import CircularReferenceError
from model_x_ray.spreadsheet.sandbox import XlsxSandbox
from model_x_ray.spreadsheet.source import XlsxSource
from model_x_ray.spreadsheet.xlsx import load_xlsx

RATE = 0.06
CASH = [180.0, 190.0, 200.0, 210.0, 220.0, 230.0]
COLUMNS = "CDEFGH"


def write_debt_schedule(path: Path, iterate: bool = True) -> Path:
    workbook = openpyxl.Workbook()
    debt = workbook.active
    debt.title = "Debt"
    debt["A1"], debt["B1"] = "Interest rate", RATE
    debt["A3"] = "EUR m"
    labels = {4: "Cash available for debt service", 5: "Opening debt", 6: "Interest on average debt", 7: "Repayment", 8: "Closing debt"}
    for row, label in labels.items():
        debt[f"A{row}"] = label
    for index, column in enumerate(COLUMNS):
        previous = COLUMNS[index - 1] if index else None
        debt[f"{column}3"] = f"{2025 + index}E"
        debt[f"{column}4"] = CASH[index]
        debt[f"{column}5"] = 1000.0 if previous is None else f"={previous}8"
        debt[f"{column}6"] = f"=$B$1*({column}5+{column}8)/2"
        debt[f"{column}7"] = f"={column}4-{column}6"
        debt[f"{column}8"] = f"={column}5-{column}7"
    workbook.calculation.iterate = iterate
    workbook.calculation.iterateCount = 100
    workbook.calculation.iterateDelta = 0.001
    workbook.save(path)
    return path


def solved_schedule() -> list[tuple[float, float]]:
    opening, rows = 1000.0, []
    for cash in CASH:
        closing = (opening * (1 + RATE / 2) - cash) / (1 - RATE / 2)
        rows.append((RATE * (opening + closing) / 2, closing))
        opening = closing
    return rows


def test_a_debt_schedule_with_interest_on_average_debt_converges_like_excel(tmp_path: Path) -> None:
    book = load_xlsx(write_debt_schedule(tmp_path / "debt.xlsx"))

    for column, (interest, closing) in zip(COLUMNS, solved_schedule()):
        assert book.value(f"Debt!{column}6") == pytest.approx(interest, abs=0.01)
        assert book.value(f"Debt!{column}8") == pytest.approx(closing, abs=0.01)

    patched = book.with_patch("Debt!B1", 0.08)
    assert patched.value("Debt!H8") > book.value("Debt!H8")


def test_with_iteration_off_the_loop_stays_a_circular_reference_as_in_excel(tmp_path: Path) -> None:
    book = load_xlsx(write_debt_schedule(tmp_path / "debt.xlsx", iterate=False))

    with pytest.raises(CircularReferenceError):
        book.value("Debt!E8")


def test_a_typed_interest_inside_the_loop_is_found_fixed_and_verified(tmp_path: Path) -> None:
    path = write_debt_schedule(tmp_path / "debt.xlsx")
    workbook = openpyxl.load_workbook(path)
    workbook["Debt"]["E6"] = 61.5
    workbook.save(path)
    setup = Setup(name="scripted", kind="investigator", model="scripted-investigator", verifier="scripted-verifier")
    source = XlsxSource({"debt": path})
    graph = build_setup_graph(setup, source, XlsxSandbox(source, tmp_path / "corrected"), "debt", lambda model_id: ScriptedVerdict() if model_id == "scripted-verifier" else ScriptedDiagnosis(turns=[([], None)]))

    final = drive_graph(graph, {"spreadsheet_id": "debt", "question": ""}, {"configurable": {"thread_id": uuid4().hex}})

    assert final["patch"]["cell"] == "Debt!E6"
    assert final["patch"]["new_formula"] == "=$B$1*(E5+E8)/2"
    assert final["verifications"][-1]["verdict"] == "confirmed"
    corrected = load_xlsx(tmp_path / "corrected" / final["corrected_copy_id"])
    assert corrected.value("Debt!H8") == pytest.approx(solved_schedule()[-1][1], abs=0.01)


def test_a_loop_that_divides_by_zero_on_the_first_pass_still_converges(tmp_path: Path) -> None:
    from model_x_ray.spreadsheet.engine import Workbook

    book = Workbook({
        "Rating!A1": "EBIT", "Rating!B1": 300.0, "Rating!A2": "Debt", "Rating!B2": 2000.0,
        "Rating!A3": "Interest coverage", "Rating!B3": "=B1/B5",
        "Rating!A4": "Spread", "Rating!B4": "=IF(B3>3,0.01,0.03)",
        "Rating!A5": "Interest expense", "Rating!B5": "=B2*(0.04+B4)",
    }, iteration=(100, 0.001))

    assert book.value("Rating!B5") == pytest.approx(140.0)
    assert book.value("Rating!B3") == pytest.approx(300 / 140)
    assert book.loop("Rating!B3")["converged"] is True
