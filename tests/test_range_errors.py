from pathlib import Path
from uuid import uuid4

import openpyxl

from conftest import ScriptedVerdict
from test_investigator import ScriptedDiagnosis

from model_x_ray.benchmark.runner import Setup, build_setup_graph, drive_graph
from model_x_ray.rules import scan_tables
from model_x_ray.spreadsheet.sandbox import XlsxSandbox
from model_x_ray.spreadsheet.source import XlsxSource
from model_x_ray.spreadsheet.xlsx import load_xlsx


def write_gas_invoice(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "jul01"
    for column, header in zip("ABCDEFG", ["Day", "North", "South", "East", "West", "Total", "Running total"]):
        sheet[f"{column}10"] = header
    for row in range(12, 39):
        sheet[f"A{row}"] = f"Jul {row - 11}"
        for column, base in zip("BCDE", [410, 275, 190, 335]):
            sheet[f"{column}{row}"] = base + (row * 7) % 23
        sheet[f"F{row}"] = f"=SUM(B{row}:E{row})"
        sheet[f"G{row}"] = f"=SUM($F$12:F{row})"
    sheet["A42"] = "TOTAL"
    for column in "BCE":
        sheet[f"{column}42"] = f"=SUM({column}12:{column}38)"
    sheet["D42"] = "=SUM(D38)"
    sheet["F42"] = "=SUM(B42:E42)"
    for column, header in zip("JKL", ["Season", "Bid average", "Mid average"]):
        sheet[f"{column}12"] = header
    for index, season in enumerate(["Summer 01", "Summer 02", "Summer 03", "Summer 04", "Summer 05"]):
        row = 13 + index
        first = 50 + 7 * index
        sheet[f"J{row}"], sheet[f"K{row}"] = season, f"=AVERAGE(F{first}:F{first + 6})"
        sheet[f"L{row}"] = f"=AVERAGE(H{first}:H{first + 6})"
    sheet["L15"] = "=AVERAGE(F64:H70)"
    sheet["F49"], sheet["G49"], sheet["H49"] = "Bid", "Ask", "Mid"
    for row in range(50, 85):
        sheet[f"F{row}"] = 3.1 + (row % 5) * 0.05
        sheet[f"G{row}"] = 3.3 + (row % 4) * 0.05
        sheet[f"H{row}"] = f"=0.4*F{row}+0.6*G{row}"
    workbook.save(path)
    return path


def write_hot_list(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Hotlist"
    sheet["A11"], sheet["K11"], sheet["M11"] = "Coal", "Deal", "Value"
    sheet["K12"], sheet["M12"] = "M&G - UK", 1000
    sheet["K20"], sheet["M20"] = "SUBTOTAL", "=SUM(M19)"
    sheet["A21"], sheet["K21"], sheet["M21"] = "Power", "Deal", "Value"
    sheet["K22"], sheet["M22"] = "Idaho Power", 200
    sheet["K23"], sheet["M23"] = "Harvard Fund", 250
    sheet["K27"], sheet["M27"] = "PSE", 300
    sheet["K29"], sheet["M29"] = "SUBTOTAL", "=SUM(M27:M28)"
    sheet["A30"], sheet["K30"], sheet["M30"] = "Gas", "Deal", "Value"
    sheet["K31"], sheet["M31"] = "Avista", 50
    sheet["K32"], sheet["M32"] = "Puget", 70
    sheet["K33"], sheet["M33"] = "SUBTOTAL", "=SUM(M31:M32)"
    sheet["K35"], sheet["M35"] = "TOTALS", "=M20+M29+M33"
    sheet["K38"], sheet["M38"] = "Deal", "Value"
    sheet["K39"], sheet["M39"] = "Spot", 40
    sheet["K42"], sheet["M42"] = "Spot", 60
    sheet["K43"], sheet["M43"] = "Largest spot deal", "=SUM(M42)"
    workbook.save(path)
    return path


def test_a_total_whose_range_disagrees_with_the_totals_beside_it_is_found_with_the_range_they_use(tmp_path: Path) -> None:
    book = load_xlsx(write_gas_invoice(tmp_path / "invoice.xlsx"))

    found = {finding["cell"]: finding for finding in scan_tables(book)}

    assert set(found) == {"jul01!D42", "jul01!L15"}
    assert (found["jul01!D42"]["kind"], found["jul01!D42"]["fix"]) == ("aggregate_range", "=SUM(D12:D38)")
    assert "C42" in found["jul01!D42"]["reason"] or "E42" in found["jul01!D42"]["reason"]
    assert (found["jul01!L15"]["kind"], found["jul01!L15"]["fix"]) == ("aggregate_range", "=AVERAGE(H64:H70)")
    patched = book.with_patch("jul01!L15", found["jul01!L15"]["fix"])
    assert patched.value("jul01!L15") != book.value("jul01!L15")


def test_a_section_subtotal_that_leaves_out_deals_above_blank_rows_is_found(tmp_path: Path) -> None:
    book = load_xlsx(write_hot_list(tmp_path / "hotlist.xlsx"))

    found = {finding["cell"]: finding for finding in scan_tables(book)}

    assert set(found) == {"Hotlist!M20", "Hotlist!M29"}
    assert (found["Hotlist!M20"]["kind"], found["Hotlist!M20"]["fix"]) == ("incomplete_total", "=SUM(M12:M19)")
    assert "M12" in found["Hotlist!M20"]["reason"]
    assert (found["Hotlist!M29"]["kind"], found["Hotlist!M29"]["fix"]) == ("incomplete_total", "=SUM(M22:M28)")
    assert book.with_patch("Hotlist!M20", "=SUM(M12:M19)").value("Hotlist!M35") == 1000 + 300 + 120


def test_the_corrected_invoice_and_hot_list_raise_no_range_findings(tmp_path: Path) -> None:
    invoice = write_gas_invoice(tmp_path / "invoice.xlsx")
    workbook = openpyxl.load_workbook(invoice)
    workbook["jul01"]["D42"], workbook["jul01"]["L15"] = "=SUM(D12:D38)", "=AVERAGE(H64:H70)"
    workbook.save(invoice)
    hot_list = write_hot_list(tmp_path / "hotlist.xlsx")
    workbook = openpyxl.load_workbook(hot_list)
    workbook["Hotlist"]["M20"], workbook["Hotlist"]["M29"] = "=SUM(M12:M19)", "=SUM(M22:M28)"
    workbook.save(hot_list)

    assert scan_tables(load_xlsx(invoice)) == []
    assert scan_tables(load_xlsx(hot_list)) == []


def test_a_hot_list_with_a_short_subtotal_is_fixed_and_verified_end_to_end(tmp_path: Path) -> None:
    path = write_hot_list(tmp_path / "hotlist.xlsx")
    setup = Setup(name="scripted", kind="investigator", model="scripted-investigator", verifier="scripted-verifier")
    source = XlsxSource({"hotlist": path})
    graph = build_setup_graph(setup, source, XlsxSandbox(source, tmp_path / "corrected"), "hotlist", lambda model_id: ScriptedVerdict() if model_id == "scripted-verifier" else ScriptedDiagnosis(turns=[([], None)]))

    final = drive_graph(graph, {"spreadsheet_id": "hotlist", "question": ""}, {"configurable": {"thread_id": uuid4().hex}})

    expected = {"M20": "=SUM(M12:M19)", "M29": "=SUM(M22:M28)"}
    target = final["target"]
    assert target["mode"] == "table" and target["cell"] in expected
    assert final["patch"]["new_formula"] == expected[target["cell"]]
    assert final["verifications"][-1]["verdict"] == "confirmed"
    assert final["verifications"][-1]["checks"]["anomaly_resolved"]
    corrected = load_xlsx(tmp_path / "corrected" / final["corrected_copy_id"])
    assert corrected.formula(f"Hotlist!{target['cell']}") == expected[target["cell"]]


def test_subtotals_over_calculations_and_totals_whose_wider_range_adds_nothing_are_not_findings() -> None:
    from model_x_ray.spreadsheet.engine import Workbook

    book = Workbook({
        "Case!G16": "Cash", "Case!H16": "Amount", "Case!H18": 1200.0, "Case!H19": 450.0, "Case!H22": "=H18-H19",
        "Case!H25": 500.0, "Case!H26": 300.0, "Case!H27": 75.0, "Case!G28": "Subtotal uses", "Case!H28": "=SUM(H25:H27)",
        "Rates!A93": "Month", "Rates!B93": "Balance", "Rates!E93": "Interest", "Rates!H93": "Fees",
        **{f"Rates!{column}{row}": 1000.0 + row for column in "BH" for row in range(94, 103)},
        **{f"Rates!E{row}": 6.5 + row / 100 for row in range(94, 102)},
        "Rates!A103": "Total", "Rates!B103": "=SUM(B94:B102)", "Rates!E103": "=SUM(E94:E101)", "Rates!H103": "=SUM(H94:H102)",
    })

    assert scan_tables(book) == []


def test_a_season_average_over_the_wrong_columns_is_fixed_and_verified_end_to_end(tmp_path: Path) -> None:
    path = write_gas_invoice(tmp_path / "invoice.xlsx")
    workbook = openpyxl.load_workbook(path)
    workbook["jul01"]["D42"] = "=SUM(D12:D38)"
    workbook.save(path)
    setup = Setup(name="scripted", kind="investigator", model="scripted-investigator", verifier="scripted-verifier")
    source = XlsxSource({"invoice": path})
    graph = build_setup_graph(setup, source, XlsxSandbox(source, tmp_path / "corrected"), "invoice", lambda model_id: ScriptedVerdict() if model_id == "scripted-verifier" else ScriptedDiagnosis(turns=[([], None)]))

    final = drive_graph(graph, {"spreadsheet_id": "invoice", "question": ""}, {"configurable": {"thread_id": uuid4().hex}})

    assert (final["target"]["mode"], final["target"]["cell"]) == ("table", "L15")
    assert final["patch"]["new_formula"] == "=AVERAGE(H64:H70)"
    assert final["verifications"][-1]["verdict"] == "confirmed"
    assert final["verifications"][-1]["checks"] == {"anomaly_resolved": True, "dependents_consistent": True, "pattern_restored": True}


def test_a_missing_dollar_sign_and_the_last_typed_input_above_carried_forward_cells_are_not_findings() -> None:
    from model_x_ray.spreadsheet.engine import Workbook

    book = Workbook({
        "Legacy!G20": "Rate", "Legacy!H20": "A", "Legacy!H21": 1.1, "Legacy!H22": 0.9, "Legacy!H23": 1.0,
        **{f"Legacy!G{row}": 100.0 + row for row in range(25, 33)},
        **{f"Legacy!H{row}": f"=$G{row}*H$21*H$22*H$23" for row in range(25, 33)},
        "Legacy!H26": "=$G26*H$21*H$22*H23",
        "Receipts!A10": "Day", "Receipts!C10": "Receipt",
        **{f"Receipts!A{row}": row - 10 for row in range(11, 31)},
        "Receipts!C11": 9000.0,
        **{f"Receipts!C{row}": f"=+C{row - 1}" for row in range(12, 15)},
        **{f"Receipts!C{row}": value for row, value in zip(range(15, 20), [10000.0, 8000.0, 10000.0, 5500.0, 0.0])},
        **{f"Receipts!C{row}": f"=+C{row - 1}" for row in range(20, 31)},
        "Receipts!C25": 4200.0,
    })

    found = {finding["cell"]: finding for finding in scan_tables(book)}

    assert set(found) == {"Receipts!C25"}
    assert (found["Receipts!C25"]["kind"], found["Receipts!C25"]["fix"]) == ("typed_in_column", "=+C24")
