from pathlib import Path

import openpyxl

from model_x_ray.rules import scan_tables
from model_x_ray.spreadsheet.xlsx import load_xlsx


def write_trading_sheets(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    deals = workbook.active
    deals.title = "Deals"
    deals.append(["Deal", "Volume", "Price", "Value", "Flag"])
    for row, (volume, price) in enumerate([(100, 2.5), (80, 2.6), (120, 2.4), (90, 2.7), (60, 2.5), (110, 2.6), (70, 2.8)], start=2):
        deals.append([f"D-{row}", volume, price, f"=B{row}*C{row}", f"=IF(B{row}>85,1,0)"])
    deals["D5"] = 288
    deals["D7"] = "=B6*C7"
    deals["E6"] = 1
    deals["A10"], deals["D10"] = "Total", "=SUM(D2:D8)"
    costs = workbook.create_sheet("Costs")
    costs.append(["Item", "Cost"])
    for row, cost in enumerate([12, 30, 7, 25, 9], start=2):
        costs.append([f"Item {row}", cost])
    costs["A8"], costs["B8"] = "Total cost", "=SUM(B2:B5)"
    costs["A9"], costs["B9"] = "Check", "=SUM(B3)"
    costs["A11"], costs["B11"] = "Fees", 4
    costs["A12"], costs["B12"] = "Taxes", 6
    costs["A13"], costs["B13"] = "Total other", "=SUM(B11:B12)"
    costs["A15"], costs["B15"] = "Grand total", "=B8+B13"
    workbook.save(path)
    return path


def test_incomplete_totals_and_column_pattern_breaks_are_found_with_their_fixes(tmp_path: Path) -> None:
    book = load_xlsx(write_trading_sheets(tmp_path / "trading.xlsx"))

    found = {finding["cell"]: finding for finding in scan_tables(book)}

    assert set(found) == {"Costs!B8", "Deals!D5", "Deals!D7", "Deals!E6"}
    assert found["Costs!B8"]["kind"] == "incomplete_total"
    assert found["Costs!B8"]["fix"] == "=SUM(B2:B6)"
    assert "B6" in found["Costs!B8"]["reason"]
    assert (found["Deals!D5"]["kind"], found["Deals!D5"]["fix"]) == ("typed_in_column", "=B5*C5")
    assert (found["Deals!D7"]["kind"], found["Deals!D7"]["fix"]) == ("column_pattern", "=B7*C7")
    assert found["Deals!E6"]["fix"] == "=IF(B6>85,1,0)"


def test_a_clean_table_has_no_findings(tmp_path: Path) -> None:
    path = write_trading_sheets(tmp_path / "trading.xlsx")
    workbook = openpyxl.load_workbook(path)
    workbook["Deals"]["D5"], workbook["Deals"]["D7"], workbook["Deals"]["E6"] = "=B5*C5", "=B7*C7", "=IF(B6>85,1,0)"
    workbook["Costs"]["B8"] = "=SUM(B2:B6)"
    workbook.save(path)

    assert scan_tables(load_xlsx(path)) == []
