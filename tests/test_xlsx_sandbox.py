from pathlib import Path

import openpyxl
import pytest
from langgraph.types import Command
from test_investigator import ScriptedDiagnosis

from model_x_ray.graph import run_investigation
from model_x_ray.spreadsheet.sandbox import XlsxSandbox, run_sandbox_test
from model_x_ray.spreadsheet.source import XlsxSource
from model_x_ray.spreadsheet.xlsx import load_xlsx


def test_an_approved_fix_is_written_to_a_new_xlsx_with_a_change_log_and_the_original_is_untouched(demo_workbook_path: Path, scripted_verdict, tmp_path: Path) -> None:
    source = XlsxSource({"demo": demo_workbook_path})
    sandbox = XlsxSandbox(source, tmp_path / "corrected")
    graph = run_investigation(ScriptedDiagnosis(turns=[(["P&L!H7"], "P&L!H7")]), "demo", source, sandbox, verifier=scripted_verdict)
    config = {"configurable": {"thread_id": "xlsx-sandbox"}}

    graph.invoke({"spreadsheet_id": "demo", "question": "Why did EBITDA 2027 jump on the P&L sheet?"}, config)
    assert graph.get_state(config).values["verifications"][-1]["after"]["P&L!H8"] == pytest.approx(44.73, abs=0.01)
    assert sandbox.open_copies() == 0

    graph.invoke(Command(resume={"action": "approve", "reason": "typed opex"}), config)
    corrected_id = graph.get_state(config).values["corrected_copy_id"]
    corrected = tmp_path / "corrected" / corrected_id

    assert corrected.exists() and corrected.suffix == ".xlsx"
    fixed = load_xlsx(corrected)
    assert fixed.formula("P&L!H7") == "=G7*(1+Assumptions!$C$7)"
    assert fixed.value("P&L!H8") == pytest.approx(44.73, abs=0.01)
    assert fixed.metadata_of("P&L!F7")["font_color"] == "000000FF"

    log = openpyxl.load_workbook(corrected)["Change log"]
    rows = [[cell.value for cell in row] for row in log.iter_rows()]
    assert rows[0] == ["Timestamp", "Cell", "Old formula", "New formula", "Reason"]
    assert rows[1][1:4] == ["P&L!H7", "-45.2", "=G7*(1+Assumptions!$C$7)"]

    original = load_xlsx(demo_workbook_path)
    assert original.raw("P&L!H7") == -45.2
    assert "Change log" not in openpyxl.load_workbook(demo_workbook_path).sheetnames


def test_a_typed_number_patch_is_written_as_a_number(demo_workbook_path: Path, tmp_path: Path) -> None:
    source = XlsxSource({"demo": demo_workbook_path})
    sandbox = XlsxSandbox(source, tmp_path / "corrected")
    name = sandbox.write_corrected_copy("demo", {"cell": "P&L!F7", "old_formula": "-47.6", "new_formula": "-48.1", "source": "human"}, [["Cell"], ["P&L!F7"]])

    assert openpyxl.load_workbook(tmp_path / "corrected" / name)["P&L"]["F7"].value == -48.1


def test_cells_that_already_fail_in_the_original_do_not_block_a_fix(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "DCF"
    sheet.append(["", "Year 1", "Year 2", "Year 3", "Year 4"])
    sheet.append(["Revenue", 100, "=B2*1.1", 121, "=D2*1.1"])
    sheet.append(["Horizon", 1, 1, 1, " "])
    sheet.append(["Cash flow", "=B2*B3", "=C2*C3", "=D2*D3", "=E2*E3"])
    workbook.save(tmp_path / "dcf.xlsx")
    source = XlsxSource({"dcf": tmp_path / "dcf.xlsx"})

    result = run_sandbox_test(source, XlsxSandbox(source, tmp_path / "corrected"), "dcf", {"cell": "DCF!D2", "old_formula": "121", "new_formula": "=C2*1.1", "source": "agent"})

    assert result["before"]["DCF!E4"] is None and result["after"]["DCF!E4"] is None
    assert result["after"]["DCF!D4"] == pytest.approx(121.0)
