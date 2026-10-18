import time
from pathlib import Path

import openpyxl
import pytest

from model_x_ray.spreadsheet.engine import CircularReferenceError, FormulaError, UnsupportedFormulaError, Workbook
from model_x_ray.spreadsheet.xlsx import load_xlsx


def test_loads_an_xlsx_model_and_recalculates_the_planted_error(demo_workbook_path: Path) -> None:
    book = load_xlsx(demo_workbook_path)

    assert book.sheets == ["P&L", "Assumptions"]
    assert book.formula("P&L!H8") == "=H6+H7"
    assert book.formula("P&L!H7") is None
    assert book.value("P&L!H7") == -45.2
    assert book.value("P&L!G8") == pytest.approx(40.03, abs=0.01)
    assert book.value("P&L!H8") == pytest.approx(52.01, abs=0.01)
    assert book.value("P&L!H9") == pytest.approx(0.299, abs=0.001)


def test_patching_a_copy_restores_the_pattern_without_touching_the_original(demo_workbook_path: Path) -> None:
    original = load_xlsx(demo_workbook_path)

    patched = original.with_patch("P&L!H7", "=G7*(1+Assumptions!$C$7)")

    assert patched.value("P&L!H7") == pytest.approx(-52.48, abs=0.01)
    assert patched.value("P&L!H8") == pytest.approx(44.73, abs=0.01)
    assert patched.value("P&L!H9") == pytest.approx(0.117, abs=0.001)
    assert original.value("P&L!H8") == pytest.approx(52.01, abs=0.01)
    assert original.formula("P&L!H7") is None


def test_dependency_graph_traces_ebitda_to_its_inputs_and_shows_the_broken_link(demo_workbook_path: Path) -> None:
    book = load_xlsx(demo_workbook_path)

    assert book.precedents("P&L!H8") == ["P&L!H6", "P&L!H7"]
    assert book.precedents("P&L!H7") == []
    assert "Assumptions!C5" in book.all_precedents("P&L!H8")
    assert "Assumptions!C7" not in book.all_precedents("P&L!H8")
    assert "Assumptions!C7" in book.all_precedents("P&L!G8")
    assert book.dependents("P&L!H7") == ["P&L!H8"]
    assert book.all_dependents("P&L!H7") == ["P&L!H8", "P&L!H9"]


def test_relative_formula_patterns_reveal_the_cell_that_breaks_its_row(demo_workbook_path: Path) -> None:
    book = load_xlsx(demo_workbook_path)

    assert book.formula_pattern("P&L!G4") == book.formula_pattern("P&L!H4")
    assert book.formula_pattern("P&L!G7") == "=RC[-1]*(1+Assumptions!R7C3)"
    assert book.formula_pattern("P&L!H7") is None
    assert book.formula_pattern("P&L!G6") == book.formula_pattern("P&L!H6") == "=R[-2]C+R[-1]C"


def test_labels_and_cell_notes_are_available_for_override_judgment(demo_workbook_with_override_note: Path) -> None:
    book = load_xlsx(demo_workbook_with_override_note)

    assert book.row_label("P&L!H7") == "Opex"
    assert book.column_header("P&L!H7") == "2027E"
    assert book.note("P&L!H7") == "Per budget meeting, opex frozen at 45.2"
    assert book.note("P&L!G7") is None


def test_supported_functions_ranges_and_errors_are_explicit() -> None:
    book = Workbook({
        "S!A1": 10.0, "S!A2": 20.0, "S!A3": 30.0,
        "S!B1": "=SUM(A1:A3)", "S!B2": "=AVERAGE(A1:A3)", "S!B3": "=MAX(A1:A3)-MIN(A1:A3)",
        "S!B4": "=IF(A1>5,ROUND(A2/3,2),0)", "S!B5": "=ABS(-A3)+2^2", "S!B6": "=50%*A2",
        "S!C1": "=C2+1", "S!C2": "=C1+1",
        "S!D1": "=INDIRECT(\"A1\")",
    })

    assert book.value("S!B1") == 60.0
    assert book.value("S!B2") == 20.0
    assert book.value("S!B3") == 20.0
    assert book.value("S!B4") == 6.67
    assert book.value("S!B5") == 34.0
    assert book.value("S!B6") == 10.0
    assert book.precedents("S!B1") == ["S!A1", "S!A2", "S!A3"]
    with pytest.raises(CircularReferenceError):
        book.value("S!C1")
    with pytest.raises(UnsupportedFormulaError, match="INDIRECT"):
        book.value("S!D1")


def test_quoted_sheet_names_link_sheets_like_excel_and_google_sheets_write_them() -> None:
    book = Workbook({
        "P&L!H8": 52.0,
        "Summary!D10": "='P&L'!H8",
        "Summary!D11": "=IF('P&L'!H8>50,\"above plan\",\"on plan\")",
    })

    assert book.precedents("Summary!D10") == ["P&L!H8"]
    assert book.value("Summary!D10") == 52.0
    assert book.value("Summary!D11") == "above plan"
    assert book.dependents("P&L!H8") == ["Summary!D10", "Summary!D11"]
    assert book.formula_pattern("Summary!D10") == "='P&L'!R[-2]C[4]"


def test_a_sheet_formatted_out_to_the_last_column_loads_only_the_cells_it_uses(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Inputs"
    sheet["B2"] = 100
    sheet["C2"] = "=B2*2"
    sheet["C2"].number_format = "0.0%"
    sheet["XFD3000"].number_format = "0.0%"
    workbook.save(tmp_path / "wide.xlsx")

    started = time.perf_counter()
    book = load_xlsx(tmp_path / "wide.xlsx")

    assert time.perf_counter() - started < 5
    assert book.value("Inputs!C2") == 200.0
    assert book.metadata_of("Inputs!C2")["number_format"] == "0.0%"
    assert book.raw("Inputs!XFD3000") is None


def test_a_chain_longer_than_the_recursion_limit_still_calculates() -> None:
    cells = {"Calc!A1": 1.0}
    for row in range(2, 4001):
        cells[f"Calc!A{row}"] = f"=A{row - 1}+1"
    book = Workbook(cells)

    assert book.value("Calc!A4000") == 4000.0
    assert book.value("Calc!A2000") == 2000.0


def test_a_link_to_another_workbook_has_no_precedents_instead_of_failing() -> None:
    book = Workbook({"Sheet!A1": "=[1]May!$J$10", "Sheet!A2": "=A1+1"})

    assert book.precedents("Sheet!A1") == []
    assert book.dependents("Sheet!A1") == ["Sheet!A2"]


def test_a_function_with_missing_arguments_is_a_formula_error_not_a_crash() -> None:
    book = Workbook({"Sheet!A1": "=EOMONTH(A2)", "Sheet!A2": 45000, "Sheet!A3": "=IFERROR(EOMONTH(A2),0)"})

    with pytest.raises(FormulaError, match="EOMONTH"):
        book.value("Sheet!A1")
    assert book.value("Sheet!A3") == 0
