import json
import random
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from model_x_ray.spreadsheet.engine import Workbook as EngineWorkbook
from model_x_ray.spreadsheet.xlsx import load_xlsx

OVERRIDE_COMMENT = "Per budget meeting: keep as entered."
PERIODS = {"F": "2024A", "G": "2025A", "H": "2026E", "I": "2027E", "J": "2028E"}
FORECAST_COLUMNS = ("H", "I", "J")
PLANT_COLUMNS = ("G", "H", "I", "J")
SINGLE_PLANT_SETS = {
    "type1": "typed_value",
    "type2": "wrong_reference",
    "type3": "inconsistent_range",
    "type4": "sign_inconsistency",
    "type5": "period_inconsistency",
}


def _previous_column(column: str) -> str | None:
    if column == "A":
        return None
    return chr(ord(column) - 1)


def _column_of(ref: str) -> str:
    return re.match(r"[A-Z]+", ref.split("!", 1)[1]).group(0)


def _row_of(ref: str) -> int:
    return int(re.search(r"\d+", ref.split("!", 1)[1]).group(0))


def _prior_ref(ref: str) -> str:
    sheet, address = ref.split("!", 1)
    column = _column_of(ref)
    previous = _previous_column(column)
    if previous is None:
        raise ValueError(f"No prior column for {ref}")
    return f"{sheet}!{previous}{address[len(column):]}"


def _pnl_cells() -> dict[str, object]:
    cells: dict[str, object] = {
        "A1": "Deal model P&L",
        "A3": "EUR m",
        "A4": "Revenue",
        "F4": 176.5,
        "G4": 188.0,
        "A5": "COGS",
        "A6": "Gross profit",
        "A7": "Direct costs",
        "A8": "Overhead",
        "A9": "Operating expenses",
        "A10": "EBITDA",
        "A11": "EBITDA margin",
    }
    for column, header in PERIODS.items():
        cells[f"{column}3"] = header
        cells[f"{column}5"] = f"=-{column}4*Assumptions!$C$5"
        cells[f"{column}6"] = f"={column}4+{column}5"
        cells[f"{column}7"] = f"=-{column}4*Assumptions!$C$6"
        cells[f"{column}8"] = f"=-{column}4*Assumptions!$C$7"
        cells[f"{column}9"] = f"=SUM({column}7:{column}8)"
        cells[f"{column}10"] = f"={column}6+{column}9"
        cells[f"{column}11"] = f"={column}10/{column}4"
        if column in FORECAST_COLUMNS:
            previous = _previous_column(column)
            cells[f"{column}4"] = f"={previous}4*(1+Assumptions!$C$4)"
    return cells


def _assumption_cells() -> dict[str, object]:
    return {
        "B4": "Revenue growth",
        "C4": 0.06,
        "B5": "COGS % of revenue",
        "C5": 0.52,
        "B6": "Direct cost % of revenue",
        "C6": 0.18,
        "B7": "Overhead % of revenue",
        "C7": 0.14,
    }


def _summary_cells() -> dict[str, object]:
    cells: dict[str, object] = {"A1": "Deal summary", "A3": "EUR m", "A5": "EBITDA", "A6": "YoY EBITDA change"}
    for column, header in PERIODS.items():
        cells[f"{column}3"] = header
        cells[f"{column}5"] = f"='P&L'!{column}10"
        if column != "F":
            previous = _previous_column(column)
            cells[f"{column}6"] = f"={column}5/{previous}5-1"
    return cells


def _base_content(ref: str) -> object:
    sheet, address = ref.split("!", 1)
    sheets = {"P&L": _pnl_cells, "Summary": _summary_cells, "Assumptions": _assumption_cells}
    return sheets[sheet]()[address]


def _write_base(path: Path) -> None:
    workbook = Workbook()
    pnl = workbook.active
    pnl.title = "P&L"
    for ref, content in _pnl_cells().items():
        pnl[ref] = content
    assumptions = workbook.create_sheet("Assumptions")
    for ref, content in _assumption_cells().items():
        assumptions[ref] = content
    summary = workbook.create_sheet("Summary")
    for ref, content in _summary_cells().items():
        summary[ref] = content
    workbook.save(path)


def _typed_sites() -> list[str]:
    return [f"P&L!{column}{row}" for column in PLANT_COLUMNS for row in (5, 6, 9, 10)] + [
        "Summary!H5",
        "Summary!I5",
        "Summary!J5",
    ]


def _wrong_reference_sites() -> list[str]:
    return (
        [f"P&L!{column}4" for column in FORECAST_COLUMNS]
        + [f"P&L!{column}5" for column in PLANT_COLUMNS]
        + ["Summary!H5", "Summary!I5", "Summary!J5"]
    )


def _range_sites() -> list[str]:
    return [f"P&L!{column}9" for column in PLANT_COLUMNS]


def _sign_sites() -> list[str]:
    return [f"P&L!{column}5" for column in PLANT_COLUMNS]


def _period_sites() -> list[str]:
    return [f"P&L!{column}4" for column in FORECAST_COLUMNS]


SITES: dict[str, list[str]] = {
    "typed_value": _typed_sites(),
    "wrong_reference": _wrong_reference_sites(),
    "inconsistent_range": _range_sites(),
    "sign_inconsistency": _sign_sites(),
    "period_inconsistency": _period_sites(),
}


def _rewire(formula: str, ref: str) -> str:
    sheet = ref.split("!", 1)[0]
    if sheet == "Summary":
        column = _column_of(ref)
        return formula.replace(f"'P&L'!{column}10", f"'P&L'!{_previous_column(column)}10")
    if _row_of(ref) == 4:
        return formula.replace("$C$4", "$C$5")
    return formula.replace("$C$5", "$C$6")


def _lag_period(formula: str, ref: str) -> str:
    previous = _previous_column(_column_of(ref))
    older = _previous_column(previous)
    return formula.replace(f"={previous}4", f"={older}4", 1)


def _plant(sheet, base_engine: EngineWorkbook, ref: str, error_type: str) -> tuple[object, str | None]:
    address = ref.split("!", 1)[1]
    original = _base_content(ref)
    comment: str | None = None
    if error_type == "typed_value":
        new_content = base_engine.value(_prior_ref(ref))
    elif error_type == "documented_override":
        new_content = base_engine.value(ref)
        comment = OVERRIDE_COMMENT
    elif error_type == "wrong_reference":
        new_content = _rewire(str(original), ref)
    elif error_type == "inconsistent_range":
        column = _column_of(ref)
        new_content = original.replace(f":{column}8)", f":{column}7)")
    elif error_type == "sign_inconsistency":
        new_content = str(original).replace("=-", "=", 1)
    else:
        new_content = _lag_period(str(original), ref)
    sheet[address] = new_content
    if comment:
        sheet[address].comment = Comment(comment, "model owner")
    return new_content, comment


def _write_variant(variant_path: Path, plants: list[tuple[str, str]], base_engine: EngineWorkbook) -> list[dict]:
    workbook = load_workbook(variant_path.parents[1] / "base.xlsx")
    records = []
    for ref, error_type in plants:
        sheet = workbook[ref.split("!", 1)[0]]
        new_content, comment = _plant(sheet, base_engine, ref, error_type)
        records.append(
            {
                "cell": ref,
                "error_type": error_type,
                "original_formula": base_engine.formula(ref),
                "new_content": new_content,
                "comment": comment,
            }
        )
    workbook.save(variant_path)
    return records


def _single_plants(rng: random.Random, set_name: str) -> list[tuple[str, str]]:
    if set_name == "clean":
        return []
    if set_name == "overrides":
        error_type = "documented_override"
    else:
        error_type = SINGLE_PLANT_SETS[set_name]
    if error_type == "documented_override":
        return [(rng.choice(SITES["typed_value"]), error_type)]
    return [(rng.choice(SITES[error_type]), error_type)]


def _blended_plants(rng: random.Random) -> list[tuple[str, str]]:
    first_pool = [
        (ref, error_type)
        for error_type in ("typed_value", "wrong_reference", "sign_inconsistency", "period_inconsistency")
        for ref in SITES[error_type]
        if ref.startswith("P&L") and _row_of(ref) in (4, 5)
    ]
    chosen = [rng.choice(first_pool)]
    count = rng.randint(2, 3)
    rest_pool = [
        (ref, error_type)
        for error_type, refs in SITES.items()
        for ref in refs
        if ref not in {cell for cell, _ in chosen} and error_type not in {kind for _, kind in chosen}
    ]
    rng.shuffle(rest_pool)
    for option in rest_pool:
        if option[0] in {cell for cell, _ in chosen}:
            continue
        chosen.append(option)
        if len(chosen) == count:
            break
    return chosen


def _usable(engine: EngineWorkbook, ref: str) -> bool:
    try:
        prior = _prior_ref(ref)
    except ValueError:
        return False
    value = engine.value(prior)
    return isinstance(value, (int, float)) and value != 0


def _target_for(engine: EngineWorkbook, planted_ref: str) -> str:
    dependents = engine.all_dependents(planted_ref)
    headline = f"P&L!{_column_of(planted_ref)}10"
    if headline in dependents and _usable(engine, headline):
        return headline
    for candidate in dependents:
        if _usable(engine, candidate):
            return candidate
    if _usable(engine, planted_ref):
        return planted_ref
    raise ValueError(f"No question target for {planted_ref}")


def _question(engine: EngineWorkbook, target: str, error_cells: list[str], verdict: str) -> dict:
    current = engine.value(target)
    previous = engine.value(_prior_ref(target))
    pct = round((current - previous) / abs(previous) * 100, 1)
    label = engine.row_label(target) or target
    year = engine.column_header(target) or ""
    return {
        "text": f"Why did {label} {year} change by {pct}%?",
        "target": target,
        "expected_error_cells": error_cells,
        "expected_verdict": verdict,
    }


def generate(output_dir: Path, per_set: int = 25, seed: int = 42) -> dict:
    rng = random.Random(seed)
    output_dir.mkdir(parents=True, exist_ok=True)
    base_path = output_dir / "base.xlsx"
    _write_base(base_path)
    base_engine = load_xlsx(base_path)
    expected_outputs = {
        ref: base_engine.value(ref) for ref in base_engine.cells() if base_engine.formula(ref) is not None
    }
    counts = {name: per_set for name in ("type1", "type2", "type3", "type4", "type5", "blended")}
    counts["overrides"] = per_set // 2
    counts["clean"] = per_set // 2
    variants = []
    for set_name, count in counts.items():
        set_dir = output_dir / set_name
        set_dir.mkdir(parents=True, exist_ok=True)
        for index in range(count):
            plants = _blended_plants(rng) if set_name == "blended" else _single_plants(rng, set_name)
            variant_path = set_dir / f"variant-{index:03d}.xlsx"
            records = _write_variant(variant_path, plants, base_engine)
            variant_engine = load_xlsx(variant_path)
            if set_name == "clean":
                questions = [
                    _question(
                        variant_engine,
                        f"P&L!{rng.choice(FORECAST_COLUMNS)}10",
                        [],
                        "no_error",
                    )
                ]
            elif set_name == "overrides":
                questions = [
                    _question(
                        variant_engine,
                        _target_for(variant_engine, records[0]["cell"]),
                        [],
                        "no_error",
                    )
                ]
            else:
                questions = [
                    _question(variant_engine, _target_for(variant_engine, record["cell"]), [record["cell"]], "root_cause")
                    for record in records
                ]
            ground_truth = {
                "set": set_name,
                "variant_index": index,
                "seed": seed,
                "planted": records,
                "questions": questions,
                "expected_outputs": expected_outputs,
            }
            ground_truth_path = set_dir / f"variant-{index:03d}.json"
            ground_truth_path.write_text(json.dumps(ground_truth, indent=2, sort_keys=True) + "\n")
            variants.append(
                {
                    "set": set_name,
                    "index": index,
                    "file": f"{set_name}/variant-{index:03d}.xlsx",
                    "ground_truth": f"{set_name}/variant-{index:03d}.json",
                }
            )
    manifest = {"seed": seed, "per_set": per_set, "counts": counts, "variants": variants}
    (output_dir / "manifest.json").write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n")
    return manifest
