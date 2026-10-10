import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Protocol
from uuid import uuid4

import openpyxl

from model_x_ray.spreadsheet.engine import Workbook, column_letters, normalize_ref, split_ref
from model_x_ray.spreadsheet.source import SpreadsheetSource, XlsxSource
from model_x_ray.state import Patch

FOLDER_MIME = "application/vnd.google-apps.folder"
SPREADSHEET_MIME = "application/vnd.google-apps.spreadsheet"
SANDBOX_FOLDER_NAME = "Model X-Ray sandbox"
CHANGE_LOG_SHEET = "Change log"
SHEETS_SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]


class Sandbox(Protocol):
    def create_copy(self, spreadsheet_id: str) -> str: ...

    def apply_patch(self, copy_id: str, patch: Patch) -> None: ...

    def read_values(self, copy_id: str, cells: list[str]) -> dict[str, object]: ...

    def discard(self, copy_id: str) -> None: ...


class SandboxWithWrite(Sandbox, Protocol):
    def write_corrected_copy(self, spreadsheet_id: str, patch: Patch, log_rows: list[list[object]]) -> str: ...


def a1_range(ref: str) -> str:
    sheet, address = split_ref(normalize_ref(ref))
    quoted = sheet if re.fullmatch(r"[A-Za-z0-9_]+", sheet) else f"'{sheet}'"
    return f"{quoted}!{column_letters(address.column)}{address.row}"


def affected_cells(book: Workbook, ref: str) -> list[str]:
    first = normalize_ref(ref)
    return [first, *book.all_dependents(first)]


def run_sandbox_test(
    source: SpreadsheetSource,
    sandbox: Sandbox,
    spreadsheet_id: str,
    patch: Patch,
) -> dict[str, dict[str, object]]:
    book = source.load(spreadsheet_id)
    cells = affected_cells(book, patch["cell"])
    before = {cell: book.value(cell) for cell in cells}
    copy_id = sandbox.create_copy(spreadsheet_id)
    try:
        sandbox.apply_patch(copy_id, patch)
        after = sandbox.read_values(copy_id, cells)
    finally:
        sandbox.discard(copy_id)
    return {"before": before, "after": after}


def build_google_services(credentials_path: str | None = None) -> tuple[Any, Any]:
    from google.oauth2 import service_account
    from googleapiclient.discovery import build

    path = credentials_path or os.environ["GOOGLE_APPLICATION_CREDENTIALS"]
    credentials = service_account.Credentials.from_service_account_file(path, scopes=SHEETS_SCOPES)
    sheets = build("sheets", "v4", credentials=credentials, cache_discovery=False)
    drive = build("drive", "v3", credentials=credentials, cache_discovery=False)
    return sheets, drive


class GoogleSandbox:
    def __init__(self, sheets: Any, drive: Any, folder_id: str | None = None) -> None:
        self._sheets = sheets
        self._drive = drive
        self._folder_id = folder_id

    def _folder(self) -> str:
        if self._folder_id is None:
            created = self._drive.files().create(
                body={"name": SANDBOX_FOLDER_NAME, "mimeType": FOLDER_MIME},
            ).execute()
            self._folder_id = created["id"]
        return self._folder_id

    def create_copy(self, spreadsheet_id: str) -> str:
        created = self._drive.files().copy(
            fileId=spreadsheet_id,
            body={"parents": [self._folder()]},
        ).execute()
        return created["id"]

    def apply_patch(self, copy_id: str, patch: Patch) -> None:
        self._sheets.spreadsheets().values().update(
            spreadsheetId=copy_id,
            range=a1_range(patch["cell"]),
            valueInputOption="USER_ENTERED",
            body={"values": [[patch["new_formula"]]]},
        ).execute()

    def read_values(self, copy_id: str, cells: list[str]) -> dict[str, object]:
        response = self._sheets.spreadsheets().values().batchGet(
            spreadsheetId=copy_id,
            ranges=[a1_range(cell) for cell in cells],
            valueRenderOption="UNFORMATTED_VALUE",
        ).execute()
        value_ranges = response.get("valueRanges", [])
        values: dict[str, object] = {}
        for position, cell in enumerate(cells):
            rows = value_ranges[position].get("values", []) if position < len(value_ranges) else []
            values[cell] = rows[0][0] if rows and rows[0] else None
        return values

    def discard(self, copy_id: str) -> None:
        self._drive.files().delete(fileId=copy_id).execute()

    def write_corrected_copy(self, spreadsheet_id: str, patch: Patch, log_rows: list[list[object]]) -> str:
        created = self._drive.files().copy(
            fileId=spreadsheet_id,
            body={"parents": [self._folder()]},
        ).execute()
        copy_id = created["id"]
        self.apply_patch(copy_id, patch)
        self._sheets.spreadsheets().batchUpdate(
            spreadsheetId=copy_id,
            body={"requests": [{"addSheet": {"properties": {"title": CHANGE_LOG_SHEET}}}]},
        ).execute()
        self._sheets.spreadsheets().values().update(
            spreadsheetId=copy_id,
            range=a1_range(f"{CHANGE_LOG_SHEET}!A1"),
            valueInputOption="USER_ENTERED",
            body={"values": log_rows},
        ).execute()
        return copy_id


def cell_content(text: str) -> float | str:
    try:
        return float(text)
    except ValueError:
        return text


class XlsxSandbox:
    def __init__(self, source: XlsxSource, output_dir: Path) -> None:
        self._source = source
        self._output_dir = output_dir
        self._copies: dict[str, Workbook] = {}

    def open_copies(self) -> int:
        return len(self._copies)

    def create_copy(self, spreadsheet_id: str) -> str:
        copy_id = uuid4().hex
        self._copies[copy_id] = self._source.load(spreadsheet_id)
        return copy_id

    def apply_patch(self, copy_id: str, patch: Patch) -> None:
        self._copies[copy_id] = self._copies[copy_id].with_patch(patch["cell"], cell_content(patch["new_formula"]))

    def read_values(self, copy_id: str, cells: list[str]) -> dict[str, object]:
        return {cell: self._copies[copy_id].value(cell) for cell in cells}

    def discard(self, copy_id: str) -> None:
        self._copies.pop(copy_id, None)

    def write_corrected_copy(self, spreadsheet_id: str, patch: Patch, log_rows: list[list[object]]) -> str:
        book = openpyxl.load_workbook(self._source.path_of(spreadsheet_id))
        sheet, address = split_ref(normalize_ref(patch["cell"]))
        book[sheet].cell(row=address.row, column=address.column).value = cell_content(patch["new_formula"])
        if CHANGE_LOG_SHEET in book.sheetnames:
            del book[CHANGE_LOG_SHEET]
        log = book.create_sheet(CHANGE_LOG_SHEET)
        for row in log_rows:
            log.append(row)
        stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
        name = f"{spreadsheet_id}-corrected-{stamp}-{uuid4().hex[:6]}.xlsx"
        self._output_dir.mkdir(parents=True, exist_ok=True)
        book.save(self._output_dir / name)
        return name
