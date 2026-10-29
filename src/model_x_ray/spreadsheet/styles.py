import colorsys
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter

THEME_ORDER = ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink")

DEFAULT_WIDTH = 8.43

BORDER_WIDTHS = {"thin": 1, "hair": 1, "dotted": 1, "dashed": 1, "medium": 2, "mediumDashed": 2, "thick": 3, "double": 3}


def theme_colors(theme: bytes | None) -> list[str]:
    if not theme:
        return []
    text = theme.decode("utf-8", "ignore")
    colors = []
    for name in THEME_ORDER:
        match = re.search(rf"<a:{name}>.*?(?:srgbClr val|lastClr)=\"([0-9A-Fa-f]{{6}})\"", text, re.S)
        colors.append(match.group(1).upper() if match else "000000")
    return colors


def tinted(rgb: str, tint: float) -> str:
    if not tint:
        return rgb
    red, green, blue = (int(rgb[index:index + 2], 16) / 255 for index in (0, 2, 4))
    hue, light, saturation = colorsys.rgb_to_hls(red, green, blue)
    light = light * (1 + tint) if tint < 0 else light * (1 - tint) + tint
    return "".join(f"{round(channel * 255):02X}" for channel in colorsys.hls_to_rgb(hue, light, saturation))


def color_of(color: Any, theme: list[str]) -> str | None:
    if color is None:
        return None
    rgb = None
    if color.type == "rgb" and isinstance(color.rgb, str) and len(color.rgb) in (6, 8):
        rgb = color.rgb[-6:]
    elif color.type == "theme" and isinstance(color.theme, int) and color.theme < len(theme):
        rgb = theme[color.theme]
    elif color.type == "indexed" and isinstance(color.indexed, int) and color.indexed < len(COLOR_INDEX):
        rgb = COLOR_INDEX[color.indexed][-6:]
    if rgb is None:
        return None
    return f"#{tinted(rgb.upper(), color.tint or 0.0)}"


def style_of(cell: Any, theme: list[str]) -> dict[str, Any]:
    style: dict[str, Any] = {}
    if cell.fill is not None and cell.fill.fill_type == "solid":
        fill = color_of(cell.fill.fgColor, theme)
        if fill:
            style["fill"] = fill
    font = cell.font
    if font is not None:
        text = color_of(font.color, theme)
        if text and text != "#000000":
            style["color"] = text
        if font.b:
            style["bold"] = True
        if font.i:
            style["italic"] = True
        if font.u:
            style["underline"] = True
        if font.sz and float(font.sz) != 11.0:
            style["size"] = float(font.sz)
    if cell.alignment is not None and cell.alignment.horizontal in ("left", "center", "right"):
        style["align"] = cell.alignment.horizontal
    if cell.number_format and cell.number_format != "General":
        style["format"] = cell.number_format
    for side in ("top", "bottom", "left", "right"):
        edge = getattr(cell.border, side, None)
        if edge is not None and edge.style:
            style[f"border_{side}"] = f"{BORDER_WIDTHS.get(edge.style, 1)}px solid {color_of(edge.color, theme) or '#000000'}"
    return style


def workbook_styles(path: Path, limits: dict[str, tuple[int, int]]) -> dict[str, dict[str, Any]]:
    book = openpyxl.load_workbook(path)
    theme = theme_colors(book.loaded_theme)
    found = {}
    for sheet in book.worksheets:
        if sheet.title not in limits:
            continue
        rows, columns = limits[sheet.title]
        table: list[dict[str, Any]] = []
        keys: dict[str, int] = {}
        cells: dict[str, int] = {}
        for row in sheet.iter_rows(min_row=1, max_row=max(rows, 1), max_col=max(columns, 1)):
            for cell in row:
                style = style_of(cell, theme)
                if not style:
                    continue
                key = repr(sorted(style.items()))
                if key not in keys:
                    keys[key] = len(table)
                    table.append(style)
                cells[cell.coordinate] = keys[key]
        widths = {}
        for index in range(1, columns + 1):
            letter = get_column_letter(index)
            dimension = sheet.column_dimensions.get(letter)
            if dimension is not None and dimension.width and abs(dimension.width - DEFAULT_WIDTH) > 0.01:
                widths[letter] = round(dimension.width * 7 + 5)
        found[sheet.title] = {"styles": table, "cell_styles": cells, "widths": widths, "gridlines": bool(sheet.sheet_view.showGridLines)}
    return found
