from datetime import date, datetime
from pathlib import Path

from openpyxl.cell.read_only import EmptyCell
from openpyxl.comments.comment_sheet import CommentSheet
from openpyxl.packaging.relationship import get_dependents, get_rels_path
from openpyxl.reader.excel import ExcelReader
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.xml.constants import COMMENTS_NS
from openpyxl.xml.functions import fromstring

from model_x_ray.spreadsheet.engine import CellContent, CellMetadata, Workbook


def to_content(value: object) -> CellContent | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, ArrayFormula):
        return value.text
    return str(value)


def font_color(cell: object) -> str | None:
    color = getattr(cell, "font", None)
    color = getattr(color, "color", None)
    if color is None or getattr(color, "type", None) != "rgb":
        return None
    return str(color.rgb)


def sheet_notes(reader: ExcelReader) -> dict[str, str]:
    notes: dict[str, str] = {}
    for sheet, rel in reader.parser.find_sheets():
        rels_path = get_rels_path(rel.target)
        if rels_path not in reader.valid_files:
            continue
        for part in get_dependents(reader.archive, rels_path).find(COMMENTS_NS):
            for ref, comment in CommentSheet.from_tree(fromstring(reader.archive.read(part.target))).comments:
                notes[f"{sheet.name}!{ref}"] = comment.text.strip()
    return notes


def load_xlsx(path: Path) -> Workbook:
    reader = ExcelReader(path, read_only=True)
    reader.read()
    source = reader.wb
    notes = sheet_notes(reader)
    names = {name: defined.attr_text for name, defined in source.defined_names.items() if defined.attr_text}
    cells: dict[str, CellContent] = {}
    metadata: dict[str, CellMetadata] = {}
    try:
        for sheet in source.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell, EmptyCell):
                        continue
                    ref = f"{sheet.title}!{cell.coordinate}"
                    content = to_content(cell.value)
                    if content is not None and content != "":
                        cells[ref] = content
                    color = font_color(cell)
                    number_format = cell.number_format
                    if color is not None or number_format != "General":
                        metadata[ref] = {"note": notes.get(ref), "font_color": color, "number_format": number_format}
    finally:
        source.close()
    return Workbook(cells, notes, source.sheetnames, metadata, names)
