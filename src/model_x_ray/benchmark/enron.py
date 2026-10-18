import random
import re
import time
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

from langchain_core.language_models import BaseChatModel
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from model_x_ray.benchmark.runner import Setup, build_setup_graph, cost_of, drive_graph, matches, percentile
from model_x_ray.spreadsheet.engine import normalize_ref
from model_x_ray.spreadsheet.sandbox import XlsxSandbox
from model_x_ray.spreadsheet.source import XlsxSource
from model_x_ray.spreadsheet.xlsx import load_xlsx

INDEX_FILE = "Enron-Errors.xlsx"
HEADER_ROWS = 5


@dataclass(frozen=True)
class EnronCase:
    error: int
    spreadsheet: int
    faulty_file: Path
    faulty_cells: list[str]
    corrected_file: Path
    corrected_cells: list[str]
    kind: str
    subtype: str
    description: str


def cells_of(sheet: str, text: object) -> list[str]:
    found = []
    for part in re.split(r"[,; ]+", str(text or "")):
        try:
            first_column, first_row, last_column, last_row = range_boundaries(part.replace("$", ""))
        except (TypeError, ValueError):
            continue
        if None in (first_column, first_row, last_column, last_row):
            continue
        found.extend(
            normalize_ref(f"{sheet}!{get_column_letter(column)}{row}")
            for row in range(first_row, last_row + 1)
            for column in range(first_column, last_column + 1)
        )
    return found


def converted(root: Path, spreadsheet: int, name: object) -> Path:
    return root / "xlsx" / str(spreadsheet) / f"{Path(str(name)).stem}.xlsx"


def read_cases(root: Path) -> list[EnronCase]:
    sheet = load_workbook(root / INDEX_FILE, read_only=True).worksheets[0]
    cases: list[EnronCase] = []
    last: list | None = None
    for row in sheet.iter_rows(min_row=HEADER_ROWS + 1, values_only=True):
        if not row or not isinstance(row[0], (int, float)):
            continue
        row = list(row) + [None] * (13 - len(row))
        if last is not None:
            for column in (1, 2, 3, 5, 6):
                row[column] = row[column] if row[column] is not None else last[column]
        last = row
        spreadsheet = int(row[1])
        cases.append(
            EnronCase(
                error=int(row[0]),
                spreadsheet=spreadsheet,
                faulty_file=converted(root, spreadsheet, row[2]),
                faulty_cells=cells_of(str(row[3]), row[4]),
                corrected_file=converted(root, spreadsheet, row[5]),
                corrected_cells=cells_of(str(row[6]), row[7]),
                kind=str(row[8] or ""),
                subtype=str(row[9] or ""),
                description=str(row[12] or ""),
            )
        )
    return cases


def split_by_spreadsheet(cases: list[EnronCase], seed: int = 42, test_share: float = 0.2) -> dict[str, list[int]]:
    errors: dict[int, int] = {}
    for case in cases:
        errors[case.spreadsheet] = errors.get(case.spreadsheet, 0) + 1
    order = sorted(errors)
    random.Random(seed).shuffle(order)
    wanted = round(test_share * len(cases))
    test: list[int] = []
    for spreadsheet in order:
        if sum(errors[s] for s in test) >= wanted:
            break
        test.append(spreadsheet)
    return {"tune": sorted(s for s in errors if s not in test), "test": sorted(test)}


def corrected_value_matches(case: EnronCase, cell: str, copy_path: Path) -> bool:
    mapped = case.corrected_cells[case.faulty_cells.index(cell)] if len(case.corrected_cells) == len(case.faulty_cells) else cell
    try:
        return matches(load_xlsx(case.corrected_file).value(mapped), load_xlsx(copy_path).value(cell))
    except Exception:
        return False


def run_case(
    setup: Setup,
    case: EnronCase,
    workdir: Path,
    corrected: bool = False,
    factory: Callable[[str], BaseChatModel] | None = None,
) -> dict:
    path = case.corrected_file if corrected else case.faulty_file
    row: dict = {"setup": setup.name, "error_nr": case.error, "spreadsheet": case.spreadsheet, "file": "corrected" if corrected else "faulty", "description": case.description, "error": None}
    if not path.is_file():
        return {**row, "error": f"missing file {path.name}"}
    spreadsheet_id = f"enron-{case.error}-{row['file']}"
    source = XlsxSource({spreadsheet_id: path})
    sandbox = XlsxSandbox(source, workdir)
    started = time.perf_counter()
    try:
        graph = build_setup_graph(setup, source, sandbox, spreadsheet_id, factory)
        final = drive_graph(graph, {"spreadsheet_id": spreadsheet_id, "question": ""}, {"configurable": {"thread_id": f"{spreadsheet_id}-{uuid4().hex[:6]}"}})
    except Exception as error:
        return {**row, "error": f"{type(error).__name__}: {str(error)[:160]}", "latency_ms": (time.perf_counter() - started) * 1000}
    latency_ms = (time.perf_counter() - started) * 1000
    verifications = final.get("verifications", [])
    patch = final.get("patch")
    confirmed = patch is not None and bool(verifications) and verifications[-1]["verdict"] == "confirmed"
    patched = normalize_ref(patch["cell"]) if patch else None
    faulty = set(case.faulty_cells)
    found = confirmed and patched in faulty
    copy_name = final.get("corrected_copy_id")
    run_log = final.get("run_log", [])
    return {
        **row,
        "root_cause_found": None if corrected else found,
        "error_seen": None if corrected else any(normalize_ref(h["cell"]) in faulty for h in final.get("hypotheses", [])),
        "fix_matches_correction": None if corrected else bool(found and copy_name and corrected_value_matches(case, patched, workdir / copy_name)),
        "false_alarm": confirmed if corrected else None,
        "patched_cell": patched,
        "verdict": verifications[-1]["verdict"] if verifications else None,
        "target": final.get("target"),
        "latency_ms": latency_ms,
        "cost_usd": cost_of(run_log),
        "input_tokens": sum(entry.get("input_tokens") or 0 for entry in run_log),
    }


def count(rows: list[dict], field: str) -> str:
    return f"{sum(1 for row in rows if row[field])} of {len(rows)}"


def summarize(rows: list[dict]) -> dict:
    ran = [row for row in rows if row["error"] is None]
    faulty = [row for row in ran if row["file"] == "faulty"]
    corrected = [row for row in ran if row["file"] == "corrected"]
    latencies = sorted(row["latency_ms"] / 1000 for row in ran)
    return {
        "cases": len(faulty),
        "root_cause_found": count(faulty, "root_cause_found"),
        "error_seen": count(faulty, "error_seen"),
        "fix_matches_correction": count(faulty, "fix_matches_correction"),
        "false_alarms": count(corrected, "false_alarm"),
        "skipped": len(rows) - len(ran),
        "latency_p50_s": round(percentile(latencies, 0.5), 1),
        "latency_p95_s": round(percentile(latencies, 0.95), 1),
        "cost_usd_total": round(sum(row["cost_usd"] for row in ran), 4),
    }
